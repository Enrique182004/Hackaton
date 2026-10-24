from __future__ import annotations

import re
import shutil
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path

from biobench_sar.rag.quiz_bank import is_answer_column, strip_document_answers

SUPPORTED_EXTENSIONS: frozenset[str] = frozenset(
    {".pdf", ".txt", ".md", ".csv", ".docx", ".pptx", ".png", ".jpg", ".jpeg"}
)
TESSERACT_HINT = (
    "Tesseract binary not found. Install it: macOS `brew install tesseract`, "
    "Ubuntu `sudo apt install tesseract-ocr`, Windows https://github.com/UB-Mannheim/tesseract/wiki"
)


# A whole document's texts (pages, slides or the one text) -> the same texts cleaned, before chunking.
Cleaner = Callable[[list[str]], list[str]]


class UnsupportedFileType(ValueError):
    pass


@dataclass(frozen=True)
class Chunk:
    text: str
    source: str
    metadata: dict[str, str | int]


def _split_long(text: str, max_chars: int) -> list[str]:
    parts: list[str] = []
    rest = text.strip()
    while len(rest) > max_chars:
        cut = max(rest.rfind(sep, 0, max_chars) for sep in (". ", "\n", " "))
        if cut < max_chars // 2:
            cut = max_chars
        else:
            cut += 1
        parts.append(rest[:cut].strip())
        rest = rest[cut:].strip()
    if rest:
        parts.append(rest)
    return parts


def chunk_paragraphs(text: str, max_chars: int = 1000) -> list[str]:
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text or "") if p.strip()]
    chunks: list[str] = []
    buf = ""
    for para in paragraphs:
        if len(para) > max_chars:
            if buf:
                chunks.append(buf)
                buf = ""
            chunks.extend(_split_long(para, max_chars))
        elif buf and len(buf) + 2 + len(para) <= max_chars:
            buf = f"{buf}\n\n{para}"
        else:
            if buf:
                chunks.append(buf)
            buf = para
    if buf:
        chunks.append(buf)
    return chunks


def _mk(texts: list[str], name: str, ftype: str, **meta: str | int) -> list[Chunk]:
    base: dict[str, str | int] = {"source": name, "type": ftype, **meta}
    return [Chunk(t.strip(), name, dict(base)) for t in texts if t and t.strip()]


def _read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def _cleaned(parts: list[str], clean: Cleaner | None) -> list[str]:
    return clean(parts) if clean is not None else parts


def _ingest_pdf(path: Path, clean: Cleaner | None = None) -> list[Chunk]:
    import pymupdf

    with pymupdf.open(str(path)) as doc:
        pages = [page.get_text() for page in doc]
    out: list[Chunk] = []
    for i, text in enumerate(_cleaned(pages, clean), start=1):
        out += _mk(chunk_paragraphs(text), path.name, "pdf", page=i)
    return out


def _ingest_csv(path: Path, max_chars: int = 1000, *, drop_answers: bool = False) -> list[Chunk]:
    import pandas as pd

    df = pd.read_csv(path, dtype=str, keep_default_na=False)
    if drop_answers:
        df = df[[col for col in df.columns if not is_answer_column(str(col))]]
    out: list[Chunk] = []
    group: list[str] = []
    first = 0

    def flush() -> None:
        if group:
            meta_row = first if len(group) == 1 else f"{first}-{first + len(group) - 1}"
            out.extend(_mk([" ".join(group)], path.name, "csv", row=meta_row))
            group.clear()

    for n, (_, row) in enumerate(df.iterrows(), start=1):
        cells = ", ".join(f"{col} is {val}" for col, val in row.items() if str(val).strip())
        line = f"Row {n}: {cells}."
        if group and len(" ".join(group)) + 1 + len(line) > max_chars:
            flush()
        if not group:
            first = n
        group.append(line)
    flush()
    return out


def _ingest_docx(path: Path, clean: Cleaner | None = None) -> list[Chunk]:
    import docx

    text = "\n\n".join(p.text for p in docx.Document(str(path)).paragraphs if p.text.strip())
    return _mk(chunk_paragraphs(_cleaned([text], clean)[0]), path.name, "docx")


def _ingest_pptx(path: Path, clean: Cleaner | None = None) -> list[Chunk]:
    from pptx import Presentation

    slides: list[str] = []
    for slide in Presentation(str(path)).slides:
        paras = [
            "".join(r.text for r in p.runs) or p.text
            for shape in slide.shapes
            if shape.has_text_frame
            for p in shape.text_frame.paragraphs
        ]
        slides.append("\n\n".join(p for p in paras if p.strip()))
    out: list[Chunk] = []
    for i, text in enumerate(_cleaned(slides, clean), start=1):
        out += _mk(chunk_paragraphs(text), path.name, "pptx", slide=i)
    return out


def _ingest_image(path: Path, clean: Cleaner | None = None) -> list[Chunk]:
    import pytesseract
    from PIL import Image

    try:
        with Image.open(path) as img:
            text = pytesseract.image_to_string(img)
    except pytesseract.TesseractNotFoundError as exc:
        raise RuntimeError(TESSERACT_HINT) from exc
    return _mk(chunk_paragraphs(_cleaned([text], clean)[0]), path.name, "image")


def ingest_file(path: str | Path, *, strip_answers: bool = False) -> list[Chunk]:
    """The file's text as chunks. strip_answers (the math profile): answer keys, "Answer:" lines, answer tables and a
    CSV's answer columns are taken out of the whole document before it is chunked, so no chunk (and no embedding)
    holds an answer. The quiz bank reads its problems and answers from the original file instead."""
    p = Path(path)
    ext = p.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileType(f"Unsupported file type: {ext or p.name}")
    clean = strip_document_answers if strip_answers else None
    if ext == ".pdf":
        return _ingest_pdf(p, clean)
    if ext in (".txt", ".md"):
        return _mk(chunk_paragraphs(_cleaned([_read_text(p)], clean)[0]), p.name, ext[1:])
    if ext == ".csv":
        return _ingest_csv(p, drop_answers=strip_answers)
    if ext == ".docx":
        return _ingest_docx(p, clean)
    if ext == ".pptx":
        return _ingest_pptx(p, clean)
    if shutil.which("tesseract") is None and not Path("/opt/homebrew/bin/tesseract").exists():
        raise RuntimeError(TESSERACT_HINT)
    return _ingest_image(p, clean)
