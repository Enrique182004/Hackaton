from __future__ import annotations

import hashlib
import os
import re
import shutil
import threading
from pathlib import Path

import pytest

from biobench_sar.rag.ingestor import (
    SUPPORTED_EXTENSIONS,
    Chunk,
    UnsupportedFileType,
    chunk_paragraphs,
    ingest_file,
)
from biobench_sar.rag.vector_store import VectorStore

DIM = 256
KNOWLEDGE = Path(__file__).resolve().parent.parent / "biobench_sar" / "knowledge"


def fake_embed(texts: list[str]) -> list[list[float]]:
    out = []
    for t in texts:
        v = [0.0] * DIM
        for w in re.findall(r"[a-z0-9]+", t.lower()):
            v[int(hashlib.md5(w.encode()).hexdigest(), 16) % DIM] += 1.0
        norm = sum(x * x for x in v) ** 0.5 or 1.0
        out.append([x / norm for x in v])
    return out


@pytest.fixture
def store(tmp_path: Path) -> VectorStore:
    return VectorStore(tmp_path / "chroma", embedder=fake_embed)


def test_chunk_paragraphs_merge_and_split() -> None:
    assert chunk_paragraphs("a\n\nb\n\n\nc", 100) == ["a\n\nb\n\nc"]
    assert chunk_paragraphs("a" * 30 + "\n\n" + "b" * 30, 40) == ["a" * 30, "b" * 30]
    long = chunk_paragraphs("word " * 100, 50)
    assert len(long) > 5 and all(len(c) <= 50 for c in long)
    assert chunk_paragraphs("  \n\n  ") == []


def test_supported_extensions() -> None:
    assert {".pdf", ".txt", ".md", ".csv", ".docx", ".pptx", ".png", ".jpg", ".jpeg"} == set(SUPPORTED_EXTENSIONS)


def test_unsupported(tmp_path: Path) -> None:
    f = tmp_path / "x.exe"
    f.write_text("hi")
    with pytest.raises(UnsupportedFileType):
        ingest_file(f)


def test_txt_md(tmp_path: Path) -> None:
    f = tmp_path / "n.md"
    f.write_text("# Title\n\nHello amoeba.", encoding="utf-8")
    chunks = ingest_file(f)
    assert chunks and chunks[0].source == "n.md" and chunks[0].metadata["type"] == "md"
    (tmp_path / "e.txt").write_text("  \n")
    assert ingest_file(tmp_path / "e.txt") == []


def test_csv(tmp_path: Path) -> None:
    f = tmp_path / "s.csv"
    f.write_text("site,soil\nA,Leptosol\nB,Feozem\n")
    chunks = ingest_file(f)
    joined = " ".join(c.text for c in chunks)
    assert "Row 1: site is A, soil is Leptosol." in joined
    assert "Row 2: site is B, soil is Feozem." in joined
    assert chunks[0].metadata["type"] == "csv" and "row" in chunks[0].metadata


def test_docx(tmp_path: Path) -> None:
    import docx

    d = docx.Document()
    d.add_paragraph("First paragraph about sterile technique.")
    d.add_paragraph("Second paragraph.")
    f = tmp_path / "a.docx"
    d.save(str(f))
    chunks = ingest_file(f)
    assert "sterile technique" in chunks[0].text and chunks[0].metadata["type"] == "docx"


def test_pptx(tmp_path: Path) -> None:
    from pptx import Presentation

    prs = Presentation()
    for title in ("Slide one text", "Slide two text"):
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = title
    f = tmp_path / "a.pptx"
    prs.save(str(f))
    chunks = ingest_file(f)
    assert [c.metadata["slide"] for c in chunks] == [1, 2]
    assert "Slide two" in chunks[1].text


def test_pdf(tmp_path: Path) -> None:
    import pymupdf

    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "Acanthamoeba cysts resist drying.")
    doc.new_page().insert_text((72, 72), "Second page content.")
    f = tmp_path / "a.pdf"
    doc.save(str(f))
    chunks = ingest_file(f)
    assert chunks[0].metadata["page"] == 1 and "Acanthamoeba" in chunks[0].text
    assert chunks[1].metadata["page"] == 2


@pytest.mark.skipif(shutil.which("tesseract") is None and not Path("/opt/homebrew/bin/tesseract").exists(),
                    reason="tesseract missing")
def test_image_ocr(tmp_path: Path) -> None:
    import pytesseract
    from PIL import Image, ImageDraw, ImageFont

    if not shutil.which("tesseract"):
        pytesseract.pytesseract.tesseract_cmd = "/opt/homebrew/bin/tesseract"
    img = Image.new("RGB", (900, 160), "white")
    try:
        font = ImageFont.truetype("/System/Library/Fonts/Helvetica.ttc", 64)
    except OSError:
        font = ImageFont.load_default(size=64)
    ImageDraw.Draw(img).text((20, 40), "AUTOCLAVE STEP", fill="black", font=font)
    f = tmp_path / "t.png"
    img.save(f)
    chunks = ingest_file(f)
    assert chunks and "autoclave" in chunks[0].text.lower()


def test_store_add_query_idempotent(store: VectorStore) -> None:
    chunks = [
        Chunk("Leptosol runoff raises turbidity", "a.md", {"type": "md"}),
        Chunk("Feozem soils are humus rich and fertile", "b.md", {"type": "md"}),
        Chunk("Autoclave all plates before disposal", "c.md", {"type": "md"}),
    ]
    assert store.add_chunks(chunks) == 3
    store.add_chunks(chunks)
    assert store.count() == 3
    res = store.query("turbidity from Leptosol runoff", k=2)
    assert len(res) == 2 and res[0].source == "a.md"
    assert res[0].score >= res[1].score and 0.0 < res[0].score <= 1.0
    assert res[0].metadata["type"] == "md"
    store.clear()
    assert store.count() == 0 and store.query("x") == []


def test_empty_and_persist(tmp_path: Path) -> None:
    s = VectorStore(tmp_path / "c", embedder=fake_embed)
    assert s.query("anything") == [] and s.add_chunks([]) == 0
    s.add_chunks([Chunk("persisted text", "p.txt", {})])
    s2 = VectorStore(tmp_path / "c", embedder=fake_embed)
    assert s2.count() == 1


def test_thread_safety(store: VectorStore) -> None:
    store.add_chunks([Chunk(f"doc {i} amoeba", "t.txt", {}) for i in range(20)])
    errors: list[Exception] = []

    def work() -> None:
        try:
            for _ in range(10):
                assert store.query("amoeba", k=3)
        except Exception as e:  # noqa: BLE001
            errors.append(e)

    ts = [threading.Thread(target=work) for _ in range(6)]
    [t.start() for t in ts]
    [t.join() for t in ts]
    assert not errors


def test_seed_knowledge_ingests_and_retrieves(store: VectorStore) -> None:
    total = 0
    for f in KNOWLEDGE.iterdir():
        if f.suffix.lower() in SUPPORTED_EXTENSIONS:
            total += store.add_chunks(ingest_file(f))
    assert total > 0
    res = store.query("Leptosol Feozem runoff turbidity", k=3)
    assert any("Leptosol" in r.text for r in res)
    res = store.query("non-nutrient agar heat-killed E. coli incubation", k=3)
    assert any("agar" in r.text.lower() for r in res)


@pytest.mark.slow
@pytest.mark.skipif(os.environ.get("SAR_SLOW") != "1", reason="set SAR_SLOW=1")
def test_real_embedder(tmp_path: Path) -> None:
    s = VectorStore(tmp_path / "c")
    s.add_chunks([Chunk("Amoebae eat bacteria", "a", {}), Chunk("Stock market rises", "b", {})])
    assert s.query("protozoa feeding on microbes", k=1)[0].source == "a"


def test_count_does_not_wait_for_model_load(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import sys
    import time
    import types

    loading = threading.Event()

    class SlowModel:
        def __init__(self, name: str) -> None:
            loading.set()
            time.sleep(1.0)

        def encode(self, texts: list[str], normalize_embeddings: bool = True) -> object:
            import numpy as np

            return np.array(fake_embed(texts))

    monkeypatch.setitem(sys.modules, "sentence_transformers", types.SimpleNamespace(SentenceTransformer=SlowModel))
    s = VectorStore(tmp_path / "c")
    worker = threading.Thread(target=s.add_chunks, args=([Chunk("alpha beta", "a.txt", {})],))
    worker.start()
    assert loading.wait(2)
    started = time.perf_counter()
    assert s.count() == 0
    assert time.perf_counter() - started < 0.1
    worker.join()
    assert s.count() == 1
    s.clear()
    assert s.count() == 0
